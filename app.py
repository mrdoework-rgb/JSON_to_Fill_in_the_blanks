import streamlit as st
import docx
import json
import io
from pathlib import Path

# Set up the page title and description
st.image("1116.jpg", use_container_width=True)

st.set_page_config(page_title="JSON to Docx Converter", page_icon="📝")
st.subheader("📝 BLANKETY BLANK Specification CLOZE task Generator")

st.markdown("FIRST use this Gemini Gem to turn a set of specification statements into a series of simple recall questions formatted as a JSON [Paste all the specification statements into this](https://gemini.google.com/gem/1_xPXVbOMLdmMlHATG7jaom9fzmSOquKK?usp=sharing).")
st.write("NEXT Paste your structured JSON below to automatically generate and download your formatted Word document.")
# Text input area for the JSON data
json_input = st.text_area(
    label="Paste Gemini's JSON Output Here:",
    height=300,
    placeholder='[\n  {\n    "topic": "Chemistry of the Atmosphere",\n    ...\n  }\n]',
)

# Button to trigger processing
if st.button("Process and Generate Document"):
    if not json_input.strip():
        st.error("Please paste some JSON data first!")
    else:
        try:
            # Parse the user-provided JSON
            data = json.loads(json_input)
            
            # Load the template document
            template_path = Path(__file__).parent / "Template.docx"
            doc = docx.Document(str(template_path))
            
            # Process the JSON arrays exactly like your target format
            for item in data:
                # 1. Append Title with "Title" style
                title_para = doc.add_paragraph(f"{item.get('topic', '')}")
                title_para.style = 'Title'
                
                # 2. Append Content Body with "normal" style
                body_para = doc.add_paragraph(f"{item.get('body_text', '')}")
                body_para.style = 'Normal'
                
                # 3. Append Isolated Word Bank with "Heading 2" style
                word_bank = item.get('word_bank', [])
                bank_string = ", ".join(word_bank)
                bank_para = doc.add_paragraph(bank_string)
                bank_para.style = 'Heading 2'
                
                # Add a blank line between sections for readability
                doc.add_paragraph("")

            # Save the document to an in-memory binary stream instead of disk
            # This allows it to run smoothly in a cloud environment
            bio = io.BytesIO()
            doc.save(bio)
            bio.seek(0)
            
            st.success("🎉 Document generated successfully!")
            
            # Create the Streamlit download button
            st.download_button(
                label="📥 Download Word Document (.docx)",
                data=bio,
                file_name="Automated_Spec_Output.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
        except json.JSONDecodeError:
            st.error("❌ Invalid JSON format! Please check that your text matches valid JSON notation (ensure quotes, commas, and brackets are correct).")
        except Exception as e:
            st.error(f"❌ An error occurred while creating the document: {e}")
